# -*- coding: utf-8 -*-
"""
Generator Dokumen Evidence & Artefak AI — Acer Smart School Awards (ASSA) 2026
Kategori: Guru Kreatif Jenjang SMP/MTs
Karya: Menyelam ke Dalam Mesin Digital
Penyusun: Ach. Chanifuddin Fanani, S.Pd. — SMP Negeri 2 Lamongan
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import re

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.abspath(os.path.join(BASE_DIR, "..", "..", ".."))

# Direktori Aset Gambar
MPI_ASSETS = os.path.join(REPO_ROOT, "lab", "mpi", "sistem-komputer", "assets")
VIDEO_ASSETS = os.path.join(REPO_ROOT, "lab", "doc", "20260930 - Apresiasi GTK 2026", "bahan_video")
THUMB_YOUTUBE = os.path.join(BASE_DIR, "thumbnail_youtube_assa2026.jpg")
IMG_VIBE_CODING = os.path.join(VIDEO_ASSETS, "02_guru_vibe_coding_16x9.jpg")
IMG_SIMULATOR = os.path.join(VIDEO_ASSETS, "03_simulator_pc_interaktif_16x9.jpg")

# Warna Resmi ASSA (Acer Teal, Slate, Muted)
PRIMARY_COLOR = RGBColor(10, 108, 88)     # Acer Teal (#0A6C58)
SECONDARY_COLOR = RGBColor(31, 41, 55)   # Slate 800 (#1F2937)
MUTED_COLOR = RGBColor(107, 114, 128)    # Slate 500 (#6B7280)
CODE_COLOR = RGBColor(55, 65, 81)        # Slate 700 (#374151)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def add_callout(doc, title, text, border_color="0A6C58", bg_color="F0FDF4"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=160, bottom=160, left=200, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    run_t = p.add_run(title)
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = PRIMARY_COLOR
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    run_b = p2.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(55, 65, 81)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_figure(doc, img_path, caption_text, width=Inches(5.4)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = MUTED_COLOR

def add_prompt_card_table(doc, items):
    """
    Menampilkan tabel katalog prompt AI dengan thumbnail gambar di kolom kiri
    dan teks prompt serta metadata di kolom kanan.
    """
    for item in items:
        img_file = item.get("img_file", "")
        img_path = os.path.join(MPI_ASSETS, img_file) if img_file else ""
        title = item.get("title", "")
        badge = item.get("badge", "")
        asset_label = item.get("asset_label", "")
        prompt_text = item.get("prompt", "")
        
        tbl = doc.add_table(rows=1, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        set_table_borders(tbl, color="E5E7EB", sz="4")
        
        cell_img = tbl.cell(0, 0)
        cell_txt = tbl.cell(0, 1)
        cell_img.width = Inches(1.3)
        cell_txt.width = Inches(5.2)
        cell_img.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        set_cell_margins(cell_img, top=100, bottom=100, left=100, right=100)
        set_cell_margins(cell_txt, top=100, bottom=100, left=120, right=120)
        set_cell_background(cell_img, "F9FAFB")
        set_cell_background(cell_txt, "FFFFFF")
        
        # Gambar Thumbnail
        p_img = cell_img.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(0)
        if os.path.exists(img_path):
            p_img.add_run().add_picture(img_path, width=Inches(1.1))
        else:
            r_noimg = p_img.add_run("[Aset Visual]")
            r_noimg.font.size = Pt(8)
            r_noimg.font.color.rgb = MUTED_COLOR
        
        # Teks Metadata & Prompt
        p_txt = cell_txt.paragraphs[0]
        p_txt.paragraph_format.space_before = Pt(0)
        p_txt.paragraph_format.space_after = Pt(2)
        
        r_t = p_txt.add_run(f"{title}  ")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = SECONDARY_COLOR
        
        r_badge = p_txt.add_run(f"[{badge}]")
        r_badge.font.name = "Calibri"
        r_badge.font.size = Pt(8)
        r_badge.font.color.rgb = PRIMARY_COLOR
        r_badge.bold = True
        
        p_file = cell_txt.add_paragraph()
        p_file.paragraph_format.space_before = Pt(0)
        p_file.paragraph_format.space_after = Pt(3)
        r_file = p_file.add_run(f"Berkas: {asset_label}")
        r_file.font.name = "Consolas"
        r_file.font.size = Pt(8)
        r_file.font.color.rgb = MUTED_COLOR
        
        p_prompt = cell_txt.add_paragraph()
        p_prompt.paragraph_format.space_before = Pt(0)
        p_prompt.paragraph_format.space_after = Pt(0)
        p_prompt.paragraph_format.line_spacing = 1.15
        r_p = p_prompt.add_run(f'"{prompt_text}"')
        r_p.font.name = "Consolas"
        r_p.font.size = Pt(8)
        r_p.font.color.rgb = CODE_COLOR
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

def generate_artefak_ai():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Acer Smart School Awards 2026 | Dokumen Evidence & Artefak AI Lengkap")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(156, 163, 175)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("SMP Negeri 2 Lamongan — Dokumentasi Prompt AI, Vibe Coding, & Atribusi Media")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(156, 163, 175)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(13.5)
        run.font.color.rgb = PRIMARY_COLOR
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = SECONDARY_COLOR
        return p

    def add_p(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            clean_prefix = bold_prefix.replace("*", "")
            r_b = p.add_run(clean_prefix)
            r_b.bold = True
            r_b.font.name = "Calibri"
            r_b.font.size = Pt(10)
            r_b.font.color.rgb = SECONDARY_COLOR
            
        tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
        for token in tokens:
            if not token:
                continue
            if token.startswith("**") and token.endswith("**") and len(token) >= 4:
                content = token[2:-2].replace("*", "")
                r = p.add_run(content)
                r.bold = True
                r.italic = italic
            elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
                content = token[1:-1].replace("*", "")
                r = p.add_run(content)
                r.italic = True
            else:
                content = token.replace("*", "")
                r = p.add_run(content)
                r.italic = italic
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(31, 41, 55)
        return p

    def add_code_block(title, code_str):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        
        tcPr = cell._element.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:left w:val="single" w:sz="18" w:space="0" w:color="0A6C58"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        rt = p.add_run(title)
        rt.bold = True
        rt.font.name = "Calibri"
        rt.font.size = Pt(9.5)
        rt.font.color.rgb = PRIMARY_COLOR
        
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.15
        rc = p2.add_run(code_str)
        rc.font.name = "Consolas"
        rc.font.size = Pt(8.5)
        rc.font.color.rgb = CODE_COLOR
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==================== HEADER JUDUL ====================
    p_kicker = doc.add_paragraph()
    p_kicker.paragraph_format.space_before = Pt(8)
    p_kicker.paragraph_format.space_after = Pt(3)
    r_kicker = p_kicker.add_run("EVIDENCE DAN ARTEFAK AI — ACER SMART SCHOOL AWARDS (ASSA) 2026")
    r_kicker.bold = True
    r_kicker.font.name = "Calibri"
    r_kicker.font.size = Pt(10)
    r_kicker.font.color.rgb = PRIMARY_COLOR

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Katalog Evidence Digital, Dokumentasi Lengkap Prompt AI, & Formulir Pengungkapan AI (AI Disclosure Statement)")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(15)
    r_title.font.color.rgb = SECONDARY_COLOR

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Inovasi: Menyelam ke Dalam Mesin Digital | Peserta: Ach. Chanifuddin Fanani, S.Pd. — SMP Negeri 2 Lamongan")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = MUTED_COLOR

    add_callout(
        doc,
        "Komitmen Kejujuran Akademik & Transparansi Radikal Pemanfaatan AI (S&K Poin 6, 9, & 10 ASSA 2026)",
        "Dokumen ini disusun untuk memberikan transparansi penuh atas seluruh aset digital dan pemanfaatan kecerdasan artifisial (AI) dalam karya inovasi 'Menyelam ke Dalam Mesin Digital'. Selaras dengan etika akademik dan ketentuan ASSA 2026, peserta membuka secara utuh: (1) Direktori live link produk dan kode sumber terbuka, (2) Formulir pengungkapan peran AI vs kendali guru, (3) Katalog lengkap 17 prompt gambar AI beserta aset acuan visualnya, (4) Dokumentasi 6-scene naskah dan prompt video AI Google Vids, (5) Engineering prompt rekayasa JavaScript/Web Audio API pada Google Antigravity, serta (6) Penegasan batas tegas komponen orisinal non-AI."
    )

    # ==================== BAGIAN 1 ====================
    add_h1("1. Direktori Bukti Implementasi & Artefak Digital (Live Links)")
    add_p("Dewan juri Acer Smart School Awards 2026 dapat menguji dan mengeksplorasi seluruh artefak digital berikut secara langsung melalui peramban web modern tanpa perlu mendaftar akun atau memasang software tambahan:")

    t_links = doc.add_table(rows=6, cols=3)
    t_links.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_links.autofit = False
    set_table_borders(t_links, color="D1D5DB", sz="6")
    
    headers = ["Nama Artefak / Produk", "Deskripsi & Fitur Kunci", "Tautan Akses Digital (Langsung)"]
    for i, h in enumerate(headers):
        cell = t_links.rows[0].cells[i]
        set_cell_background(cell, "0A6C58")
        set_cell_margins(cell, top=90, bottom=90, left=80, right=80)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); r.bold = True; r.font.name = "Calibri"; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(255, 255, 255)

    link_data = [
        ("Media Pembelajaran Interaktif (MPI): Sistem Komputer", "21 Halaman modular lanskap 16:9, simulator perakitan PC, alur booting, indikator BIOS beep, evaluasi formatif terpadu, dan dokumentasi prompt AI.", "https://fanani.my.id/lab/mpi/sistem-komputer/"),
        ("Laboratorium Maya: Perakitan Komputer", "Lab simulasi drag-and-drop hardware ATX motherboard, tes inkompatibilitas hardware, gerbang logika dasar, sakelar biner 8-bit, dan LKPD digital cetak PDF.", "https://fanani.my.id/lab/lab-maya/perakitan-komputer/index.html"),
        ("Repositori Kode Sumber Terbuka (GitHub)", "Repositori kode program sumber terbuka (HTML5/CSS3/Vanilla JS) lengkap dengan commit history pengembangan dan dokumentasi teknis.", "https://github.com/chanif/chanif.github.io"),
        ("Video Narasi Praktik Baik (YouTube)", "Video dokumentasi implementasi di ruang kelas 9A dan laboratorium komputer SMPN 2 Lamongan (Durasi 04:48 Menit, Resolusi Full HD, Takarir Indonesia).", "https://www.youtube.com/watch?v=p-auk5ULNFo"),
        ("Dokumentasi Diseminasi & Pengimbasan", "Materi tayang, dokumentasi foto, dan modul pelatihan bagi 160 guru SMP se-Kabupaten Lamongan bersama Dinas Pendidikan Kabupaten Lamongan.", "https://fanani.my.id/edu/workshop/2026-09-workshop-kka-dinas-pendidikan/index.html")
    ]

    for row_idx, (name, desc, url) in enumerate(link_data, start=1):
        row = t_links.rows[row_idx]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(2.1); c1.width = Inches(2.4); c2.width = Inches(2.0)
        set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c2, top=60, bottom=60, left=80, right=80)
        if row_idx % 2 == 1:
            set_cell_background(c0, "F9FAFB"); set_cell_background(c1, "F9FAFB"); set_cell_background(c2, "F9FAFB")
            
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(name); r0.font.name = "Calibri"; r0.font.size = Pt(8.5); r0.bold = True; r0.font.color.rgb = SECONDARY_COLOR
        
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(desc); r1.font.name = "Calibri"; r1.font.size = Pt(8); r1.font.color.rgb = RGBColor(75, 85, 99)
        
        p2 = c2.paragraphs[0]; p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(url); r2.font.name = "Calibri"; r2.font.size = Pt(8); r2.font.color.rgb = PRIMARY_COLOR; r2.font.underline = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_figure(doc, IMG_SIMULATOR, "Gambar 1: Antarmuka Laboratorium Maya Perakitan Komputer (Live Demo: https://fanani.my.id/lab/lab-maya/perakitan-komputer/).", width=Inches(5.4))

    # ==================== BAGIAN 2 ====================
    add_h1("2. Formulir Pengungkapan Pemanfaatan AI (AI Disclosure Statement)")
    add_p("Berdasarkan Ketentuan Angka 10 Regulasi Acer Smart School Awards 2026, berikut matriks batasan penggunaan kecerdasan buatan dan peran kontrol manusia (guru):")

    t_ai = doc.add_table(rows=6, cols=3)
    t_ai.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_ai.autofit = False
    set_table_borders(t_ai, color="D1D5DB", sz="6")
    
    ai_headers = ["Alat AI yang Digunakan", "Tujuan & Bagian yang Dibantu AI", "Peran & Kontrol Verifikasi Guru"]
    for i, h in enumerate(ai_headers):
        cell = t_ai.rows[0].cells[i]
        set_cell_background(cell, "0A6C58")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); r.bold = True; r.font.name = "Calibri"; r.font.size = Pt(9); r.font.color.rgb = RGBColor(255, 255, 255)

    ai_data = [
        ("Google Antigravity (AI Coding Assistant)", "Membantu akselerasi penulisan algoritma interaktif JavaScript, kalibrasi fungsi deteksi batas koordinat drag-and-drop, dan optimasi CSS.", "Guru menetapkan arsitektur sistem, skenario logika inkompatibilitas komponen (misal CPU LGA vs Socket AM4), serta menguji fungsionalitas di panel sentuh IFP kelas."),
        ("Google Vids / AI Video Generator", "Membantu pembuatan draf animasi visualisasi siklus mesin Von Neumann (Fetch, Decode, Execute, Store) beresolusi 1080p.", "Guru menyusun storyboard alur kerja bus CPU, memeriksa akurasi konsep ilmiah komputer, dan menyunting durasi agar sesuai rentang atensi siswa SMP."),
        ("Image Generation AI (DALL-E / Imagen)", "Membantu pembuatan ilustrasi metafora komputer (misalnya analogi motherboard dan komponen pendukung hardware).", "Guru memfilter prompt dan gambar agar bebas bias, ramah anak, serta menempatkan atribusi dan prompt lengkap di halaman 'Dokumentasi Prompt AI' media."),
        ("LLM / Asisten Teks AI", "Membantu penyusunan variasi distraktor (pilihan jawaban pengecoh) pada bank soal formatif pilihan ganda dan menjodohkan.", "Guru memvalidasi butir soal sesuai Capaian Pembelajaran (CP) Fase D, mengkalibrasi tingkat kesukaran HOTS, dan memastikan kebenaran kunci jawaban."),
        ("Prinsip Integritas Data (Zero AI Falsification)", "AI TIDAK DIGUNAKAN untuk memalsukan data dampak, nilai siswa, refleksi siswa, testimoni orang tua, atau dokumentasi kelas.", "Seluruh data kuantitatif (rata-rata nilai 58,40 -> 86,80; ketuntasan 91,4%) dan foto kegiatan adalah data riil kelas IX-A SMP Negeri 2 Lamongan.")
    ]

    for row_idx, (tool, purpose, control) in enumerate(ai_data, start=1):
        row = t_ai.rows[row_idx]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(2.0); c1.width = Inches(2.3); c2.width = Inches(2.2)
        set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c2, top=60, bottom=60, left=80, right=80)
        if row_idx % 2 == 1:
            set_cell_background(c0, "F9FAFB"); set_cell_background(c1, "F9FAFB"); set_cell_background(c2, "F9FAFB")
            
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(tool); r0.font.name = "Calibri"; r0.font.size = Pt(8.5); r0.bold = True; r0.font.color.rgb = SECONDARY_COLOR
        
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(purpose); r1.font.name = "Calibri"; r1.font.size = Pt(8); r1.font.color.rgb = RGBColor(75, 85, 99)
        
        p2 = c2.paragraphs[0]; p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(control); r2.font.name = "Calibri"; r2.font.size = Pt(8); r2.font.color.rgb = RGBColor(55, 65, 81)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ==================== BAGIAN 3 ====================
    add_h1("3. Katalog Lengkap Prompt Gambar AI (Midjourney / DALL-E / Imagen)")
    add_p("Seluruh ilustrasi visual yang ditampilkan pada aplikasi MPI Sistem Komputer didokumentasikan secara transparan beserta deskripsi teks prompt orisinal yang digunakan:")

    add_h2("3.1 Prompt Gambar AI — Menu Utama & Latar Belakang (7 Aset Visual)")
    
    prompts_menu = [
        {
            "img_file": "bg-lab-clean.jpg",
            "title": "Latar Belakang Utama MPI (Lab Komputer Modern & Cerah)",
            "badge": "Background Visual MPI",
            "asset_label": "assets/bg-lab-clean.jpg",
            "prompt": "Clean minimalist 3D empty modern computer laboratory and digital workspace background. Bright spacious interior with large sunlit glass windows overlooking green trees, clean sleek white and pastel cyan desks with neat minimalist monitors, smooth light wood laminate floor, soft ambient morning daylight, peaceful spacious and airy, soft pastel color palette of clean white, soft teal, light cream, and warm sunlight. Completely empty room with no people, minimal and serene, 16:9 aspect ratio."
        },
        {
            "img_file": "thumb_materi_1.jpg",
            "title": "Sistem Komputer & Segitiga Emas",
            "badge": "Sub-Materi 1",
            "asset_label": "assets/thumb_materi_1.jpg",
            "prompt": "Clean minimalist 3D isometric illustration representing the Computer System Golden Triangle. A glowing stylized triangle linking three clean icons: a friendly desktop computer monitor representing Hardware, a holographic glowing software window representing Software, and a cute friendly student avatar icon representing Brainware. Solid uniform soft pastel teal mint background (#b2dfdb), cute modern educational illustration, centered composition, 16:9 aspect ratio."
        },
        {
            "img_file": "thumb_materi_2.jpg",
            "title": "Tantangan Menjodohkan Hardware",
            "badge": "Sub-Materi 2",
            "asset_label": "assets/thumb_materi_2.jpg",
            "prompt": "Clean minimalist 3D illustration representing an interactive hardware matching challenge and puzzle game. Cute stylized computer hardware components like a mechanical keyboard, a gaming mouse, and a RAM memory stick connecting with vibrant puzzle connector pieces or target matching slots. Solid uniform soft pastel yellow amber background (#fff9c4), cheerful playful educational design, clean render, 16:9 aspect ratio."
        },
        {
            "img_file": "thumb_materi_3.jpg",
            "title": "Perangkat Keras & Motherboard ATX",
            "badge": "Sub-Materi 3",
            "asset_label": "assets/thumb_materi_3.jpg",
            "prompt": "Clean minimalist 3D isometric illustration of a computer motherboard and core hardware components. A cute stylized ATX motherboard with an open golden CPU socket in the center, colorful RAM slots, heat sinks, and a PCIe graphics card slot, surrounded neatly by simplified storage SSD and processor chip. Solid uniform soft pastel sky blue background (#bbdefb), modern educational tech aesthetic, clean render, 16:9 aspect ratio."
        },
        {
            "img_file": "thumb_materi_4.jpg",
            "title": "Simulasi Video Siklus CPU",
            "badge": "Sub-Materi 4",
            "asset_label": "assets/thumb_materi_4.jpg",
            "prompt": "Clean minimalist 3D illustration of a multimedia video player screen presenting an animated CPU data simulation. A cute stylized tablet monitor screen displaying glowing circuit data waves and a large glossy circular video play button icon in the center. Solid uniform soft pastel rose pink background (#f8bbd0), modern educational aesthetic, clean render, 16:9 aspect ratio."
        },
        {
            "img_file": "thumb_materi_5.jpg",
            "title": "Cara Kerja Komputer (Fetch-Decode-Execute)",
            "badge": "Sub-Materi 5",
            "asset_label": "assets/thumb_materi_5.jpg",
            "prompt": "Clean minimalist 3D isometric illustration representing the internal CPU machine cycle (Fetch, Decode, Execute, Store). A glowing central CPU chip with miniature clockwork gears and 4 colorful glowing curved circular arrows looping in a continuous cycle, with glowing binary 1s and 0s and small RAM memory chips beside it. Solid uniform soft pastel orange peach background (#ffe0b2), modern clean educational tech style, 16:9 aspect ratio."
        },
        {
            "img_file": "thumb_materi_6.jpg",
            "title": "Sistem Operasi & Manajemen Kolaborasi",
            "badge": "Sub-Materi 6",
            "asset_label": "assets/thumb_materi_6.jpg",
            "prompt": "Clean minimalist 3D isometric illustration of an Operating System desktop interface connecting hardware and apps. Cute floating layers of modern graphical user interface windows (GUI), colorful app tiles, a gears settings icon, and a multitasking process dashboard floating harmoniously above a digital platform. Solid uniform soft pastel lavender purple background (#e1bee7), modern educational tech aesthetic, 16:9 aspect ratio."
        }
    ]
    add_prompt_card_table(doc, prompts_menu)

    add_h2("3.2 Prompt Gambar AI — Aset Simulator PC & Kasus Diagnostik (10 Aset Visual)")
    
    prompts_sim = [
        {
            "img_file": "comp_cpu_good.jpg",
            "title": "CPU Multi-Core LGA 1700 (Komponen Tepat)",
            "badge": "Misi 1: Rakit PC",
            "asset_label": "assets/comp_cpu_good.jpg",
            "prompt": "Realistic modern desktop CPU microprocessor chip with LGA 1700 gold contact pads on bottom, shiny metallic integrated heat spreader (IHS) on top, clean tech product studio shot, crisp macro detail, soft cyan studio lighting, white isolated background."
        },
        {
            "img_file": "comp_cpu_bad.jpg",
            "title": "CPU Jadul Pin Patah / AM3 (Pengecoh Rusak)",
            "badge": "Misi 1: Distraktor",
            "asset_label": "assets/comp_cpu_bad.jpg",
            "prompt": "Old vintage computer CPU processor chip with several bent and broken golden pins on green circuit substrate, visibly damaged electronic component, macro product photography, studio lighting, isolated neutral background."
        },
        {
            "img_file": "comp_ram_good.jpg",
            "title": "RAM 16GB DDR4 3200MHz (Komponen Tepat)",
            "badge": "Misi 1: Rakit PC",
            "asset_label": "assets/comp_ram_good.jpg",
            "prompt": "Modern high performance desktop RAM stick with sleek black metal heat spreader and subtle RGB light diffuser, gold connector pins with DDR4 notch alignment, clean studio product shot, light blue background."
        },
        {
            "img_file": "comp_ram_bad.jpg",
            "title": "RAM 2GB DDR1/DDR2 Lawas (Pengecoh Inkompatibel)",
            "badge": "Misi 1: Distraktor",
            "asset_label": "assets/comp_ram_bad.jpg",
            "prompt": "Vintage legacy DDR1 SDRAM memory stick module with exposed green printed circuit board and retro memory chips, aged outdated computer hardware with legacy notch position, clean studio product photography."
        },
        {
            "img_file": "comp_ssd_good.jpg",
            "title": "SSD NVMe M.2 PCIe 512GB (Komponen Tepat)",
            "badge": "Misi 1: Rakit PC",
            "asset_label": "assets/comp_ssd_good.jpg",
            "prompt": "Modern M.2 NVMe solid state drive SSD with gold contact pins, flash memory NAND chips, and controller, ultra high speed storage component, clean macro product photography, soft studio lighting."
        },
        {
            "img_file": "comp_floppy_bad.jpg",
            "title": "Disket Floppy Disk 1.44MB (Pengecoh Lawas)",
            "badge": "Misi 1: Distraktor",
            "asset_label": "assets/comp_floppy_bad.jpg",
            "prompt": "Vintage 3.5 inch black magnetic floppy diskette with metal sliding shutter, retro obsolete storage media from 1990s, clean studio lighting, isolated neutral background."
        },
        {
            "img_file": "comp_psu_good.jpg",
            "title": "PSU 550W 80+ Bronze ATX (Komponen Tepat)",
            "badge": "Misi 1: Rakit PC",
            "asset_label": "assets/comp_psu_good.jpg",
            "prompt": "Modern 550W ATX desktop power supply unit (PSU) with black metal casing, honeycomb ventilation grill, cooling fan, and bundled power cables, clean studio shot, neutral background."
        },
        {
            "img_file": "comp_charger_bad.jpg",
            "title": "Adaptor Daya Handphone 10W DC (Pengecoh Kurang Daya)",
            "badge": "Misi 1: Distraktor",
            "asset_label": "assets/comp_charger_bad.jpg",
            "prompt": "Small white 10W USB mobile phone wall charger plug adapter with two metal prongs, isometric 3D render, clean isolated light background, tech icon for low power unsuitable adapter."
        },
        {
            "img_file": "trouble_screen_beep.jpg",
            "title": "Layar Blank NO SIGNAL & Beep Alert (Kasus Diagnostik 1)",
            "badge": "Misi 3: Troubleshooting",
            "asset_label": "assets/trouble_screen_beep.jpg",
            "prompt": "Desktop computer monitor displaying a black screen with a floating No Signal warning message, connected to a modern PC tower with power LED on, alert sound waves icon, clean isometric 3D render, educational computer lab troubleshooting illustration."
        },
        {
            "img_file": "trouble_task_manager.jpg",
            "title": "Task Manager RAM 98% Overload (Kasus Diagnostik 2)",
            "badge": "Misi 4: Troubleshooting",
            "asset_label": "assets/trouble_task_manager.jpg",
            "prompt": "Computer desktop screen displaying a diagnostic Task Manager performance window, prominent red 98 percent RAM memory overload bar chart, clean modern UI graphics, isometric 3D render, computer troubleshooting educational illustration."
        }
    ]
    add_prompt_card_table(doc, prompts_sim)

    # ==================== BAGIAN 4 ====================
    add_h1("4. Transparansi Komponen Non-AI & Orisinalitas Vektor Interaktif (SVG)")
    add_p("Untuk menjamin akurasi ilmiah dan presisi teknis, bagian-bagian kunci berikut **TIDAK MENGGUNAKAN AI GRAFIS**, melainkan disusun murni secara manual oleh guru:")
    add_p(" Papan sirkuit Motherboard ATX, soket LGA 1700 dengan tuas pengunci, lekukan notch konektor RAM DDR4, pinout M.2 NVMe, slot ekspansi PCIe x16, dan jalur bus data tembaga dirancang 100% menggunakan kode vektor SVG (Scalable Vector Graphics). Pendekatan vektor manual ini esensial agar area drop target memiliki presisi piksel yang akurat dan tidak pecah saat diperbesar di layar panel sentuh IFP kelas.", bold_prefix="1. Anatomi Motherboard Interaktif (Vektor SVG Murni):")
    add_p(" Seluruh logo resmi yang disematkan—meliputi Logo Kementerian Pendidikan Dasar dan Menengah (Kemendikdasmen RI), Logo Sobat SMP 2025, Logo Gerakan Pendidikan Bermutu, Logo Sekolah Ramah Anak, serta foto profil pendidik—merupakan materi identitas resmi institusional non-AI.", bold_prefix="2. Logo Lembaga dan Atribusi Identitas Resmi:")

    # ==================== BAGIAN 5 ====================
    add_h1("5. Dokumentasi Lengkap Produksi Video Animasi AI (Google Vids)")
    add_p("Animasi simulasi siklus mesin CPU (Fetch-Decode-Execute-Store) diproduksi melalui platform Google Vids (Google Workspace AI) dengan aset keluaran 'assets/video/siklus_cpu.mp4'. Berikut rincian alur kerja dan prompt pengembangannya:")

    add_h2("5.1 Alur Kerja Produksi Video di Google Vids")
    add_p(" Guru memasukkan deskripsi skenario siklus CPU ke fitur AI Google Vids untuk menghasilkan garis besar babak cerita, pemenggalan scene otomatis, serta draf narasi edukatif.", bold_prefix="1. Prompt Storyboard & Naskah:")
    add_p(" Mengaitkan gambar acuan 3D beresolusi Full HD (16:9) pada setiap segmen scene (Klik Mouse, RAM Fetch, CPU Decode, ALU Execute, GPU Store, Monitor Output).", bold_prefix="2. Penyiapan Aset Visual Scene:")
    add_p(" Menggunakan suara AI Voiceover bawaan Google Vids untuk membacakan penjelasan alur kerja siklus CPU secara jernih, runtut, dan ramah bagi siswa SMP.", bold_prefix="3. Penerapan Narasi Suara AI (AI Voiceover):")
    add_p(" Mengatur transisi antar adegan pada timeline Google Vids, kemudian mengekspor video beresolusi 1080p Full HD (16:9).", bold_prefix="4. Finishing Timeline & Ekspor:")

    add_h2("5.2 Master Prompt & Storyboard Narasi (Bilingual)")
    add_code_block(
        "Master Prompt Storyboard Google Vids (Versi Bahasa Inggris):",
        "Create an engaging, modern 3D educational video explaining the internal computer machine cycle for junior high school students. The video follows a student clicking a mouse, zooming smoothly into the computer case to reveal the motherboard and processor. Show how data is retrieved from RAM (Fetch), decoded inside the Control Unit (Decode), calculated within the Arithmetic Logic Unit (Execute), and stored back before sending output to the graphics card and monitor screen (Store & Output). Include clear, inspiring educational narration, sleek tech visuals with cyan and gold circuit glow, clean motion graphics, and a professional closing scene. 16:9 widescreen, Full HD quality."
    )
    add_code_block(
        "Master Prompt Storyboard Google Vids (Versi Bahasa Indonesia):",
        "Buat video pembelajaran 3D yang memukau dan edukatif tentang siklus kerja internal komputer (siklus mesin CPU) untuk siswa SMP. Video diawali dari siswa berseragam sekolah yang mengklik mouse di meja belajarnya, lalu kamera menyelam ke dalam casing PC memperlihatkan papan motherboard ATX berbalut sirkuit emas dan toska (#00ACC1). Tampilkan modul RAM melepaskan instruksi (Fetch), Control Unit menguraikan kode (Decode), ALU mengeksekusi perhitungan logika matematika (Execute), dan hasilnya disimpan serta dialirkan ke GPU hingga muncul di layar monitor (Store & Output). Sertakan narasi suara yang jelas dan bersahabat, visual teknologi modern beresolusi tinggi, dan penutupan yang inspiratif. Format 16:9 Full HD."
    )

    add_h2("5.3 Rangkaian Prompt Visual 6-Scene Timeline (Durasi 48 Detik)")
    
    scenes = [
        ("Scene 1 (00:00 – 00:08): Dari Klik Mouse Menyelam ke Jalur Motherboard (Input Trigger)",
         "Continuing from the start frame. The student's finger clicks the mouse button. Instantly, an electric cyan data pulse travels down the cable and injects into the motherboard USB header. The camera performs a smooth, cinematic zoom push-in, passing through the transparent glass panel and diving into a microscopic view of the ATX motherboard. Golden and cyan data pulses race through intricate copper PCB traces toward the central processor socket. Fluid 60fps motion, dynamic camera glide, neon circuitry glow, clean educational aesthetic, no text clutter."),
        ("Scene 2 (00:08 – 00:16): Penjemputan Instruksi dari RAM (Tahap 1: FETCH)",
         "Continuing along the motherboard surface. The camera pans smoothly across an illuminated DDR4 RAM module with glowing memory heat spreaders. From within the memory cells, brilliant glowing data packets representing binary instruction codes float out. A high-speed system data bus lights up in vivid cyan as the instructions are retrieved (Fetched) and rush at lightning speed toward the massive central CPU processor socket. Cinematic motion blur, glowing data streams, crisp depth of field."),
        ("Scene 3 (00:16 – 00:24): Penguraian Kode di Control Unit (Tahap 2: DECODE)",
         "Continuing into the microscopic core of the CPU microprocessor. The camera glides into a futuristic silicon chamber filled with crystalline semiconductor logic gates. At the center, the Control Unit (CU)—a glowing hexagonal control hub—scans the incoming binary packet with a soft scanning laser. The binary code seamlessly unfolds (Decodes) into distinct colored control micro-signals. Synchronization indicators pulse in perfect harmony with the processor's high-speed clock cycle rhythm."),
        ("Scene 4 (00:24 – 00:32): Komputasi Kilat di ALU (Tahap 3: EXECUTE)",
         "Continuing into the Arithmetic Logic Unit (ALU) chamber. The control signals enter a dynamic computational matrix. Mathematical and logic symbols (+, ×, AND, OR) pulse inside glowing energy spheres. The ALU executes the calculation instantly; two binary data streams collide at a focal point, merging into a brilliant golden spark that signifies a successfully computed result within nanoseconds. High-tech dynamic lighting, crisp particle effects."),
        ("Scene 5 (00:32 – 00:40): Penyimpanan & Saluran ke GPU (Tahap 4: STORE & DISPLAY BUS)",
         "Continuing from the computed result. The output is written back (Stored) to high-speed internal registers, while graphical output streams rapidly along the glowing golden PCIe x16 expansion lanes. The camera tracks alongside the data surge as it enters the Graphics Card (GPU). The dual cooling fans spin gracefully, metallic heatsink fins shimmer with soft RGB illumination, and the GPU transforms the raw data into ultra-high-resolution pixel color packets."),
        ("Scene 6 (00:40 – 00:48): Tampil Seketika di Layar Monitor (OUTPUT TUNTAS)",
         "Continuing from the GPU. The pixel packets stream through the display port and burst onto the monitor's display panel. The camera smoothly and dramatically zooms out through the microscopic pixel matrix back into the warm study room. On the large computer monitor, a vibrant interactive 3D simulation renders instantly and flawlessly. The student smiles in awe, realizing how trillions of microscopic machine cycles execute in the blink of an eye. Inspiring classroom ambient glow, cinematic closing sequence.")
    ]

    for sc_title, sc_prompt in scenes:
        add_code_block(sc_title, sc_prompt)

    # ==================== BAGIAN 6 ====================
    add_h1("6. Engineering Prompt & Arsitektur Vibe Coding (Google Antigravity)")
    add_p("Pendekatan Vibe Coding memposisikan guru sebagai Arsitek Utama Perangkat Lunak (*Lead Software Architect & Pedagogue*) yang berdialog intensif dengan asisten AI Google Antigravity untuk merekayasa logika JavaScript dan CSS native. Berikut adalah cuplikan prompt rekayasa perangkat lunak utama yang digunakan:")

    add_code_block(
        "Prompt 1: Kalibrasi Pointer Events & Responsivitas Multitouch IFP",
        "Rancang modul kontrol drag-and-drop JavaScript berbasis Web Pointer Events terpadu (pointerdown, pointermove, pointerup, setPointerCapture). Fungsi harus menangani input mouse desktop dan sentuhan multi-touch pada layar besar Interactive Flat Panel (IFP) tanpa konflik scroll halaman (touch-action: none). Lakukan normalisasi kalkulasi koordinat offset getBoundingClientRect() agar posisi ujung jari pada layar sentuh presisi 1:1 dengan titik penempatan komponen."
    )

    add_code_block(
        "Prompt 2: Deteksi Tabrakan Koordinat (Bounding Box Collision) & Logika Anti-Distraktor",
        "Buat algoritma validasi area target perakitan komponen PC. Ketika elemen komponen di-drop, periksa apakah titik pusat elemen berada di dalam area bounding box soket target motherboard (misal soket CPU LGA 1700). Validasi tipe komponen: jika komponen sesuai (CPU LGA 1700), kunci posisi elemen tepat di tengah soket, aktifkan status terpasang, dan bunyikan audio beep sukses. Jika komponen adalah pengecoh (misal CPU AM3 pin patah atau RAM DDR1), tolak pemasangan, kembalikan elemen ke rak komponen dengan transisi CSS kubik halus, dan munculkan toast error diagnostik yang edukatif bagi murid SMP."
    )

    add_code_block(
        "Prompt 3: Pembangkitan Audio BIOS Beep Berbasis Web Audio API (Zero External MP3)",
        "Tuliskan modul audio mandiri tanpa ketergantungan file MP3 eksternal menggunakan Web Audio API murni. Buat fungsi playBiosBeep(type) yang memanfaatkan AudioContext, OscillatorNode (tipe gelombang sinus 'sine'), dan GainNode. Untuk booting normal sukses, bangkitkan 1 beep pendek frekuensi 800Hz durasi 0.12 detik dengan envelope gain cepat. Untuk kasus kegagalan hardware (trouble RAM/VGA), bangkitkan rangkaian beep panjang berulang frekuensi 400Hz durasi 0.4 detik yang dipisahkan jeda 0.2 detik."
    )

    add_code_block(
        "Prompt 4: Mesin Evaluasi Formatif dengan Garis Penghubung Dinamis SVG (Menjodohkan)",
        "Implementasikan modul latihan menjodohkan interaktif 2-kolom (Komponen Hardware vs Fungsi Komputasi). Saat murid mengklik butir di kolom kiri dan menyeret ke kolom kanan, render elemen garis dinamis SVG (<line> atau <path>) secara real-time yang menghubungkan titik tengah kedua elemen. Saat koneksi dilepas, validasi pasangan: ubah warna garis menjadi hijau jika benar dan merah jika salah, serta hitung akumulasi skor formatif secara otomatis."
    )

    add_figure(doc, IMG_VIBE_CODING, "Gambar 2: Dokumentasi praktik Vibe Coding bersama AI — Menjaga transparansi peran guru sebagai arsitek pedagogis dan AI sebagai co-developer.", width=Inches(5.4))

    # ==================== BAGIAN 7 ====================
    add_h1("7. Dokumentasi Video Praktik Baik di YouTube")
    add_p("Video narasi praktik baik implementasi pembelajaran berdurasi 04:48 menit telah dipublikasikan di kanal YouTube pribadi dengan status Publik, beresolusi Full HD 1080p, dan dilengkapi Takarir (Closed Captions) Bahasa Indonesia:")
    add_p(" https://www.youtube.com/watch?v=p-auk5ULNFo", bold_prefix="Tautan Video YouTube:")

    add_figure(doc, THUMB_YOUTUBE, "Gambar 3: Cuplikan video dokumentasi praktik baik pembelajaran berdurasi 04:48 menit pada kanal YouTube (Takarir Indonesia).", width=Inches(5.4))

    # ==================== BAGIAN 8 ====================
    add_h1("8. Pernyataan Orisinalitas, Lisensi Terbuka, dan HKI")
    add_p("1. Karya inovasi 'Menyelam ke Dalam Mesin Digital' adalah karya orisinal peserta yang dikembangkan bersama asisten AI secara etis, transparan, dan bertanggung jawab untuk kepentingan pembelajaran murid di SMP Negeri 2 Lamongan.")
    add_p("2. Seluruh aset grafis SVG, diagram arsitektur komputer, struktur logika kode program JavaScript/HTML5, dan lembar kerja peserta didik merupakan hasil ciptaan mandiri guru yang dilisensikan secara terbuka (Open Source) bagi kemajuan pendidikan Indonesia.")
    add_p("3. Peserta memberikan hak penuh dan izin non-eksklusif kepada PT Acer Indonesia untuk mendokumentasikan, menampilkan, dan mempublikasikan artefak digital ini bagi kepentingan penyelenggaraan Acer Smart School Awards 2026.")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Tanda Tangan
    p_sign_date = doc.add_paragraph()
    p_sign_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign_date.paragraph_format.space_after = Pt(2)
    r_sd = p_sign_date.add_run("Lamongan, 26 Oktober 2026\nYang membuat pernyataan & pengungkapan AI,")
    r_sd.font.name = "Calibri"
    r_sd.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(36)

    p_sign_name = doc.add_paragraph()
    p_sign_name.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign_name.paragraph_format.space_after = Pt(0)
    r_sn = p_sign_name.add_run("ACH. CHANIFUDDIN FANANI, S.Pd.\n")
    r_sn.bold = True
    r_sn.font.name = "Calibri"
    r_sn.font.size = Pt(11)
    r_sn2 = p_sign_name.add_run("NIP. 19910826 202012 1 007 | NUPTK. 9158 7696 7013 0183\nGuru Informatika SMP Negeri 2 Lamongan")
    r_sn2.font.name = "Calibri"
    r_sn2.font.size = Pt(9.5)
    r_sn2.font.color.rgb = MUTED_COLOR

    out_path = r"d:\Belajar\vibecoding\chanif.github.io\lab\doc\20261026 - Acer Smart School Awards 2026\Guru Kreatif_Ach Chanifuddin Fanani_Menyelam Mesin Digital_Artefak AI.docx"
    doc.save(out_path)
    print("Artefak AI document saved successfully to:", out_path)

if __name__ == "__main__":
    generate_artefak_ai()
